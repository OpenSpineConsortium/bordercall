"""Five slides on naming a vertebra exactly, for a CNS talk or a NeurIPS oral, from one builder.

    python slides/make_exact_slides.py                      # both decks into slides/exact/
    python slides/make_exact_slides.py --venue neurips      # one of them
    python slides/make_exact_slides.py --refresh-equations  # set the equations again with pandoc
    python slides/make_exact_slides.py --public             # no confidentiality marking (see --help)

The owner's ask of 2026-10-26 20:45 UTC: figures and five PowerPoint slides, with the figures,
the text and the LaTeX as separate elements, styled like Benzel with inspiration from Kevin
Murphy's chapter on Bayes' rule, ready for the Congress of Neurological Surgeons or a NeurIPS
main-conference oral. One builder makes both decks from the same record:

  neurips   a plain 16:9 deck, white, the slide's whole width, with the method: every equation,
            the likelihood, the covariance, the selective guarantee and the claims by summing
  cns       the CNS 2026 speaker template, read from the junction deck this folder already holds
            (slides/junction_slides.pptx, built on the owner's template), inside the template's
            band and with its navy titles; results only, since the talk shows what the rule does
            and not how (docs/project.md): no equation, and no method word on a slide, in alt text,
            in a note or in the file's properties

Every slide is a title, pictures, text boxes and, at NeurIPS, equations, each its own element and
named so in the selection pane: a figure is one picture per panel (exact_slide_figures.py, drawn
for each venue at the exact inches its deck gives it, so the deck never scales a picture, its
caption as the picture's alternative text); every word is a text box; every equation is native
PowerPoint mathematics (Office Math, editable in PowerPoint's equation editor) set from its LaTeX
by pandoc, with the LaTeX in the element's alternative text, and a picture of the same equation
drawn by matplotlib as the fallback an application without Office Math shows. pandoc runs only
when an equation is new or changed, or with --refresh-equations: its output is kept beside the
deck (slides/exact/equations/), so a build needs no pandoc. Every number, and every constant in an
equation, is read from records/exact-2026-10-26/exact.json and supplement.json, never typed.

Both decks carry a confidentiality marking in the footer until they are built with --public, which
is for use only after the provisional is filed and the TTO has read the deck. The build checks the
saved file, every XML part of it: neither deck may hold a word of the confidential work
(make_junction_slides.FORBIDDEN) or a template author, and the CNS deck may hold no method word
(CNS_REFUSED); a deck that fails is deleted and the build stops.
"""
from __future__ import annotations

import argparse
import copy
import datetime
import decimal
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from typing import Dict, List, Optional, Sequence, Tuple, Union

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, "bordercall"))
sys.path.insert(0, HERE)
from make_junction_slides import FORBIDDEN, NAVY      # noqa: E402
import exact_slide_figures as XS                      # noqa: E402

RECORD = os.path.join(ROOT, "records", "exact-2026-10-26")
OUT = os.path.join(HERE, "exact")
EQDIR = os.path.join(OUT, "equations")
CNS_TEMPLATE = os.path.join(HERE, "junction_slides.pptx")
INK, BLUE, GREY, RED = (RGBColor(0x11, 0x11, 0x11), RGBColor(0x1B, 0x4F, 0x8C), RGBColor(0x59, 0x59, 0x59),
                        RGBColor(0xA4, 0x26, 0x2C))
SERIF = "Cambria"
MATH = "Cambria Math"
SPEAKER = "Gregory Schwing"
BYLINE = ("Gregory Schwing, MD, PhD",
          "General Surgery (Wayne State University School of Medicine /\vDetroit Medical Center)")
SOURCE = ("Synthetic columns from Panjabi's published tables (Spine 1991;16:861, 888; 1992;17:299); "
          "no patient data.")
MARKING = "Confidential: pre-filing draft, not for distribution"
BUILT = datetime.datetime(2026, 10, 26)
# Words of the method the results-only CNS deck may not hold anywhere in the saved file. The
# anatomical "posterior body height" (and "posterior height") is a measure, not the method.
CNS_REFUSED = re.compile(
    r"bayes|likelihood|\bposterior\b(?!\s+(?:body\s+)?height)|\bprior\b|gauss|hermite|quadrature|rank-one|"
    r"sherman|covariance|enumerat|generator|13/12|1/12|latex|equation", re.IGNORECASE)
# Authors a template brings with it: the CNS template's and python-pptx's defaults.
TEMPLATE_AUTHORS = ("gbur", "canny", "python-pptx")
NS = {"m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
      "a14": "http://schemas.microsoft.com/office/drawing/2010/main"}
M_, A_ = "{%s}" % NS["m"], "{%s}" % NS["a"]
EMU = 914400.0

Para = Union[str, Sequence[Tuple[str, dict]]]
I = {"italic": True}


# ---- the record ------------------------------------------------------------------------------------

class Numbers:
    """Every number a slide or a note says, read from the record, with its formatting."""

    def __init__(self, record: str):
        with open(os.path.join(record, "exact.json"), encoding="utf-8") as fh:
            self.d = json.load(fh)
        with open(os.path.join(record, "supplement.json"), encoding="utf-8") as fh:
            self.s = json.load(fh)
        import exact_figures as EF
        self.EF = EF

    def both(self, n: int) -> dict:
        return self.d["widths"][str(n)]["regimes"]["both"]

    def exact(self, n: int) -> float:
        return self.both(n)["exact"]

    def by_centre(self, n: int, level: str = "T12") -> float:
        return self.both(n)["by_centre"][level]["exact"]

    def sure(self, n: int) -> dict:
        return self.both(n)["selective"]["0.99"]

    def least_sure_precision(self) -> Tuple[float, int, int]:
        """The lowest precision at 0.99 over every width the record scores, and the widths' range."""
        widths = sorted(int(n) for n in self.d["widths"])
        return min(self.sure(n)["precision"] for n in widths), widths[0], widths[-1]

    def classifier(self, n: int) -> float:
        return self.d["classifier"]["both"][str(n)]["exact"]

    def gap_range(self) -> Tuple[float, float]:
        """How far the linear classifier falls below the rule, in points, on the same first 1,000
        columns it was scored on, over every width it was scored at."""
        gaps = [100 * (self.s["first_columns"][n]["first"] - self.d["classifier"]["both"][n]["exact"])
                for n in self.d["classifier"]["both"]]
        return min(gaps), max(gaps)

    def first_columns(self) -> int:
        """How many columns the linear classifier was scored on: the first ones of the record's."""
        one = self.s["first_columns"]["1"]
        return int(one["fields_first"]) // self.levels

    def top_bin(self, n: int) -> dict:
        """The calibration bin of highest posterior that holds a field, with its errors counted."""
        c = [b for b in self.both(n)["calibration"] if b["n"] > 0][-1]
        return {"n": int(c["n"]), "errors": int(round(c["n"] * (1.0 - c["correct"]))), "lo": c["lo"]}

    def challenger(self, name: str, n: int = 3) -> float:
        return self.d["challengers"][str(n)][name]["exact"]

    def cut(self, n: int, name: str) -> dict:
        return self.both(n)["claims"]["cuts"][name]

    def regions(self, n: int) -> dict:
        return self.both(n)["claims"]["regions"]

    def region_far(self, n: int) -> dict:
        return self.s["region_split"]["generator"][str(n)]["far"]

    def world(self, key: str, rule: str, n: int = 3) -> dict:
        path = next(w[1] for w in self.EF.WORLDS if w[0] == key)
        return self.EF.world_numbers(self.EF.world_cells(self.d["robust"], path, rule)[str(n)])

    def called(self, rule: str, region: str, n: int = 3, world: str = "table_specimens") -> dict:
        return self.d["robust"][world][rule][str(n)]["selective_099_by_called_region"][region]

    @property
    def columns(self) -> int:
        return int(self.d["n_columns"])

    @property
    def levels(self) -> int:
        return len(self.d["levels"])

    @property
    def measures(self) -> int:
        return len(self.d["measures"])

    @property
    def nodes(self) -> int:
        return int(self.d["gh_nodes"])

    @property
    def tau(self) -> float:
        return float(self.d["synth"]["scale_sd"])

    @property
    def rho(self) -> float:
        return float(self.d["synth"]["rho"])

    @property
    def noise_mm(self) -> float:
        return float(self.d["synth"]["noise_mm"])

    @property
    def honest(self) -> dict:
        return self.d["robust"]["rules"]["honest"]

    @property
    def table_n(self) -> int:
        return int(self.honest["table_n"])

    @property
    def honest_mm(self) -> float:
        return float(self.honest["noise_mm"])

    @property
    def robust_columns(self) -> Tuple[int, int, int]:
        r = self.d["robust"]
        return int(r["misspecified_columns"]), int(r["worlds"]), int(r["columns_per_world"])


def pct(x: float, digits: int = 0) -> str:
    """A share as a percentage, rounded half up from the number as the record stores it, so 0.945
    reads 95% and not the 94% its binary double would round to."""
    q = decimal.Decimal(1).scaleb(-digits)
    return "%s%%" % (decimal.Decimal(repr(x)) * 100).quantize(q, rounding=decimal.ROUND_HALF_UP)


def count(k: int) -> str:
    return "{:,}".format(int(k))


def equations(N: Numbers) -> Dict[str, str]:
    """Every equation on the NeurIPS slides, as LaTeX, by the name its element carries, its
    constants from the record. Each is set twice from this one source: by pandoc as Office Math,
    and by matplotlib's mathtext as the fallback picture, so it uses only what both read."""
    h = N.honest
    return {
        "problem": (r"x_{1:n} = (x_1, \ldots, x_n),\quad x_i \in \mathbb{R}^{%d},\qquad k \in \{1, \ldots, %d-n\}"
                    % (N.measures, N.levels + 1)),
        "bayes": r"p(k \mid x_{1:n}) = \dfrac{p(x_{1:n} \mid k)\, p(k)}{\sum_{k'}\; p(x_{1:n} \mid k')\, p(k')}",
        "likelihood": (r"p(x_{1:n} \mid k) = \int \prod_{m=1}^{%d} \mathcal{N}\left(x_{m,1:n} \mid s\,\mu_{m,k:k+n-1},\;"
                       r"\Sigma_{m,k:k+n-1}(s)\right) p(s)\, ds" % N.measures),
        "covariance": (r"\Sigma_{m,k:k+n-1}(s) = s^2 \rho^2\, \sigma \sigma^{\mathrm{T}} + "
                       r"\mathrm{diag}\left(s^2 (1-\rho^2)\, \sigma^2 + \sigma_{\mathrm{err}}^2\right)"),
        "size": (r"\log s \sim \mathcal{N}(0,\, %g^2),\qquad \rho = %g,\qquad \sigma_{\mathrm{err}} = %g\ \mathrm{mm}"
                 % (N.tau, N.rho, N.noise_mm)),
        "selective": (r"\Pr\left(\hat{k} = k \mid \hat{p} \geq 1-\varepsilon\right) = "
                      r"\mathbb{E}\left[\hat{p} \mid \hat{p} \geq 1-\varepsilon\right] \geq 1-\varepsilon"),
        "estimate": r"\hat{k} = \arg\max_{k}\, p(k \mid x_{1:n}),\qquad \hat{p} = p(\hat{k} \mid x_{1:n})",
        "middle": r"\ell = k + \lfloor (n-1)/2 \rfloor",
        "claim": r"\Pr(\ell \geq c \mid x_{1:n}) = \sum_{k:\, \ell \geq c} p(k \mid x_{1:n})",
        "honest": (r"\tilde{\sigma}_{m,\ell} = \sigma_{m,\ell}\sqrt{1 + 1/%d},\qquad \sigma_{\mathrm{err}} = %g\ \mathrm{mm}"
                   % (h["table_n"], h["noise_mm"])),
    }


# ---- equations ----------------------------------------------------------------------------------------

def cached_omml(name: str, tex: str, refresh: bool) -> etree._Element:
    """The equation as Office Math (m:oMathPara), from the cache beside the deck when it was set from
    this same LaTeX, else from pandoc, which converts LaTeX to Office Math in a Word file."""
    path = os.path.join(EQDIR, name + ".omml.xml")
    if not refresh and os.path.exists(path):
        root = etree.parse(path).getroot()
        if root.get("tex") == tex:
            return root[0]
    if not shutil.which("pandoc"):
        raise SystemExit("equation %r is new or changed and pandoc is not installed to set it" % name)
    d = tempfile.mkdtemp()
    with open(os.path.join(d, "e.md"), "w", encoding="utf-8") as fh:
        fh.write("$$" + tex + "$$\n")
    subprocess.run(["pandoc", os.path.join(d, "e.md"), "-o", os.path.join(d, "e.docx")], check=True)
    doc = etree.fromstring(zipfile.ZipFile(os.path.join(d, "e.docx")).read("word/document.xml"))
    shutil.rmtree(d, ignore_errors=True)
    para = doc.find(".//m:oMathPara", NS)
    if para is None:
        raise SystemExit("pandoc gave no equation for %r" % name)
    keep = etree.Element("equation", tex=tex)
    keep.append(copy.deepcopy(para))
    os.makedirs(EQDIR, exist_ok=True)
    etree.ElementTree(keep).write(path, encoding="utf-8", xml_declaration=True, pretty_print=True)
    return para


def fallback_png(name: str, tex: str, size_pt: float, eqdir: str = EQDIR) -> str:
    """The same equation as a picture at the size it is set, 300 dpi, for an application with no
    Office Math (LibreOffice among them). It is drawn afresh on every build, into `eqdir`, while the
    Office Math is read from the cache beside the committed deck."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams["mathtext.fontset"] = "cm"
    os.makedirs(eqdir, exist_ok=True)
    path = os.path.join(eqdir, "%s.png" % name)
    fig = plt.figure(figsize=(0.01, 0.01))
    fig.text(0, 0, "$" + tex + "$", fontsize=size_pt, color="#111111")
    fig.savefig(path, dpi=300, bbox_inches="tight", pad_inches=0.03, transparent=True)
    plt.close(fig)
    return path


def styled_math(para: etree._Element, size_pt: float, colour: str = "111111") -> etree._Element:
    """Office Math runs in PowerPoint carry DrawingML run properties: Cambria Math at the size the
    slide sets, italic as mathematics is."""
    para = copy.deepcopy(para)
    for jc in para.iter(M_ + "jc"):
        jc.set(M_ + "val", "left")
    for r in para.iter(M_ + "r"):
        for old in r.findall(A_ + "rPr"):
            r.remove(old)
        rpr = etree.Element(A_ + "rPr", lang="en-US", i="1", sz=str(int(round(size_pt * 100))))
        fill = etree.SubElement(rpr, A_ + "solidFill")
        etree.SubElement(fill, A_ + "srgbClr", val=colour)
        etree.SubElement(rpr, A_ + "latin", typeface=MATH)
        etree.SubElement(rpr, A_ + "cs", typeface=MATH)
        t = r.find(M_ + "t")
        (t.addprevious if t is not None else r.append)(rpr)
    return para


def attr(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


def add_equation(slide, name: str, tex: str, x: Emu, y: Emu, size_pt: float, refresh: bool,
                 max_w: Optional[Emu] = None, eqdir: str = EQDIR):
    """The equation as one element: Office Math in a text box, with its picture as the fallback, the
    LaTeX in its alternative text. The box takes the picture's size, scaled down to `max_w`."""
    png = fallback_png(name, tex, size_pt, eqdir)
    pic = slide.shapes.add_picture(png, x, y)
    if max_w is not None and pic.width > max_w:
        k = max_w / pic.width
        size_pt *= k
        pic.height = Emu(int(pic.height * k))
        pic.width = max_w
    cx, cy, rid = pic.width, pic.height, pic._element.blipFill.blip.rEmbed
    tree = pic._element.getparent()
    at = tree.index(pic._element)
    tree.remove(pic._element)
    sid = 1 + max(int(e.get("id")) for e in tree.iter("{%s}cNvPr" % NS["p"]))
    descr = attr("LaTeX: " + tex)
    xml = (
        '<mc:AlternateContent xmlns:mc="{mc}" xmlns:p="{p}" xmlns:a="{a}" xmlns:r="{r}">'
        '<mc:Choice xmlns:a14="{a14}" Requires="a14"><p:sp><p:nvSpPr>'
        '<p:cNvPr id="{sid}" name="Equation: {name}" descr="{descr}"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
        '<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
        '<p:txBody><a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" anchor="ctr"><a:noAutofit/></a:bodyPr>'
        '<a:lstStyle/><a:p><a:pPr algn="l"/><a14:m/><a:endParaRPr lang="en-US" sz="{sz}"/></a:p></p:txBody>'
        '</p:sp></mc:Choice><mc:Fallback><p:sp><p:nvSpPr>'
        '<p:cNvPr id="{sid}" name="Equation: {name}" descr="{descr}"/>'
        '<p:cNvSpPr txBox="1"><a:spLocks noRot="1" noChangeAspect="1" noMove="1" noResize="1" noEditPoints="1" '
        'noAdjustHandles="1" noChangeArrowheads="1" noChangeShapeType="1" noTextEdit="1"/></p:cNvSpPr><p:nvPr/></p:nvSpPr>'
        '<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/>'
        '</a:stretch></a:blipFill></p:spPr><p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US"><a:noFill/>'
        '</a:rPr><a:t> </a:t></a:r></a:p></p:txBody></p:sp></mc:Fallback></mc:AlternateContent>'
    ).format(sid=sid, name=name, descr=descr, x=int(x), y=int(y), cx=int(cx), cy=int(cy), rid=rid,
             sz=int(round(size_pt * 100)), **NS)
    el = etree.fromstring(xml)
    el.find(".//{%s}m" % NS["a14"]).append(styled_math(cached_omml(name, tex, refresh), size_pt))
    tree.insert(at, el)
    return cx, cy


# ---- the deck ------------------------------------------------------------------------------------

class Venue:
    """Where a deck puts things: the content box's left edge and the scale from the plain deck's
    inches, the title's colour, the template and the layout."""

    def __init__(self, name: str):
        self.name = name
        self.cns = name == "cns"
        if self.cns:
            self.template, self.layout = CNS_TEMPLATE, 2
            self.x0, self.scale, self.title_colour = 1.914, XS.CNS_SCALE, NAVY
        else:
            self.template, self.layout = None, 5
            self.x0, self.scale, self.title_colour = 0.60, 1.0, INK

    def x(self, inches: float) -> Emu:
        """A left edge given in the plain deck's inches, where the content box starts at 0.6."""
        return Inches(self.x0 + (inches - 0.60) * self.scale)

    def w(self, inches: float) -> Emu:
        return Inches(inches * self.scale)


class Deck:
    def __init__(self, venue: Venue, figures: str, refresh: bool, eqdir: str = EQDIR,
                 eq: Optional[Dict[str, str]] = None, public: bool = False):
        self.v, self.figures, self.refresh, self.eqdir = venue, figures, refresh, eqdir
        self.eq, self.public = eq or {}, public
        if venue.template:
            self.prs = Presentation(venue.template)
            lst = self.prs.slides._sldIdLst
            for sid in list(lst):
                self.prs.part.drop_rel(sid.rId)
                lst.remove(sid)
        else:
            self.prs = Presentation()
            self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.layout = self.prs.slide_layouts[venue.layout]
        self.words: List[str] = []
        self.count = 0
        self.title = ""

    def slide(self, title: str):
        """A slide with the layout's own title box, set where the template sets it and anchored at its
        foot so a second line grows upward, every other placeholder removed."""
        s = self.prs.slides.add_slide(self.layout)
        self.count += 1
        self.title = self.title or title
        t = s.shapes.title
        t.left, t.top, t.width, t.height = self.v.x(0.6), Inches(0.30), self.v.w(12.13), Inches(1.05)
        t.text_frame.text = title
        t.text_frame.word_wrap = True
        t.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
        for para in t.text_frame.paragraphs:
            para.alignment = PP_ALIGN.LEFT
            for run in para.runs:
                run.font.size, run.font.bold, run.font.name = Pt(28), True, SERIF
                run.font.color.rgb = self.v.title_colour
        for ph in list(s.placeholders):
            if ph.placeholder_format.idx != 0:
                ph._element.getparent().remove(ph._element)
        self.words.append(title)
        return s

    def text(self, s, name: str, x: float, y: float, w: float, h: float, paras: Sequence[Para], size: float = 16,
             colour: RGBColor = INK, align=PP_ALIGN.LEFT, after: float = 6, anchor=MSO_ANCHOR.TOP, bold=False,
             line: Optional[float] = None):
        """A text box at the plain deck's inches; a vertical tab in a run is a line break."""
        tb = s.shapes.add_textbox(self.v.x(x), Inches(y), self.v.w(w), Inches(h))
        tb.name = "Text: " + name
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.02)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.space_after = Pt(after)
            if line:
                p.line_spacing = line
            runs = [(para, {})] if isinstance(para, str) else para
            for text, o in runs:
                for j, piece in enumerate(text.split("\v")):
                    if j:
                        p.add_line_break()
                    r = p.add_run()
                    r.text = piece
                    r.font.name = SERIF
                    r.font.size = Pt(o.get("size", size))
                    r.font.bold = o.get("bold", bold)
                    r.font.italic = o.get("italic", False)
                    r.font.color.rgb = o.get("colour", colour)
                self.words.append(text)
        return tb

    def stat(self, s, name: str, x: float, y: float, w: float, big: str, small: Para, colour: RGBColor = BLUE,
             big_size: float = 36, small_size: float = 14, h: float = 1.15):
        """A large number with its sentence beneath: one text box, two paragraphs."""
        small_runs = [(small, {"size": small_size})] if isinstance(small, str) else \
            [(t, dict(o, size=small_size)) for t, o in small]
        return self.text(s, "stat " + name, x, y, w, h, [[(big, {"size": big_size, "bold": True, "colour": colour})],
                                                          small_runs], after=1, line=0.95)

    def figure(self, s, name: str, x: float, y: float):
        """One panel at the exact inches it was drawn at for this venue, never scaled, its caption
        as alt text."""
        w = XS.SIZES[self.v.name][name][0]
        pic = s.shapes.add_picture(os.path.join(self.figures, name + ".png"), self.v.x(x), Inches(y), width=Inches(w))
        pic.name = "Figure: " + name
        with open(os.path.join(self.figures, name + ".caption.txt"), encoding="utf-8") as fh:
            caption = fh.read().strip()
        pic._element.nvPicPr.cNvPr.set("descr", caption)
        self.words.append(caption)
        return pic

    def equation(self, s, name: str, x: float, y: float, size: float,
                 max_w: Optional[float] = None) -> Tuple[float, float]:
        """The equation at the plain deck's inches; its width and height in inches."""
        if self.v.cns:
            raise SystemExit("the CNS deck shows results only and takes no equation (%s)" % name)
        tex = self.eq[name]
        self.words.append(tex)
        cx, cy = add_equation(s, name, tex, self.v.x(x), Inches(y), size, self.refresh,
                              None if max_w is None else self.v.w(max_w), self.eqdir)
        return cx / EMU, cy / EMU

    def footer(self, s):
        """Where the numbers come from, and until the filing the confidentiality marking."""
        paras = [SOURCE] + ([] if self.public else [[(MARKING, {"bold": True, "colour": INK})]])
        self.text(s, "footer", 0.6, 6.86, 10.9, 0.45, paras, size=11, colour=GREY, after=0)
        self.text(s, "slide number", 11.73, 6.86, 1.0, 0.3, ["%d / 5" % self.count], size=11, colour=GREY,
                  align=PP_ALIGN.RIGHT, after=0)

    def notes(self, s, text: str):
        s.notes_slide.notes_text_frame.text = text
        self.words.append(text)

    def check(self):
        """No word of the confidential work, and never the banned word for where data came from."""
        low = " ".join(self.words).lower()
        bad = [w for w in FORBIDDEN + ("prov" + "enance",) if w in low]
        if bad:
            raise SystemExit("refusing the deck: %s on a slide or in a note" % ", ".join(bad))

    def properties(self):
        """The file's own properties name the speaker, never a template's author, and carry the
        build's fixed date."""
        cp = self.prs.core_properties
        cp.author = cp.last_modified_by = SPEAKER
        cp.title = self.title
        cp.subject = cp.keywords = cp.category = cp.comments = ""
        cp.revision = 1
        cp.created = cp.modified = BUILT


APP_XML = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
           '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
           'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
           '<TotalTime>0</TotalTime><Slides>%d</Slides><Notes>%d</Notes><PresentationFormat>Widescreen'
           '</PresentationFormat></Properties>')


def rewrite_app_xml(path: str, slides: int):
    """The extended properties a template brings name its own slides and fonts: write them anew."""
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = (APP_XML % (slides, slides)).encode("utf-8") if item.filename == "docProps/app.xml" \
                else src.read(item.filename)
            dst.writestr(item, data)
    os.replace(tmp, path)


def xml_parts(path: str) -> Dict[str, str]:
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n).decode("utf-8", "replace") for n in z.namelist()
                if n.endswith(".xml") or n.endswith(".rels")}


def refusals(path: str, cns: bool) -> List[Tuple[str, str]]:
    """What the saved deck may not hold, part by part: a word of the confidential work, the banned
    word, a template author, and in the CNS deck a word of the method."""
    bad = []
    for name, text in xml_parts(path).items():
        low = text.lower()
        for w in FORBIDDEN + ("prov" + "enance",) + TEMPLATE_AUTHORS:
            if w in low:
                bad.append((name, w))
        if cns:
            bad += [(name, m.group(0)) for m in CNS_REFUSED.finditer(text)]
    return bad


# ---- the five slides -------------------------------------------------------------------------------

def slide_problem(deck: Deck, N: Numbers, cns: bool):
    s = deck.slide("Can the bodies in view name a vertebra?")
    deck.figure(s, "exact_slide_view", 0.6, 1.45)
    deck.figure(s, "exact_slide_measures", 3.0, 1.55)
    x, w = 7.05, 5.68
    deck.text(s, "byline", x, 1.45, w, 0.85, [[(BYLINE[0], {"bold": True})], BYLINE[1]], size=14, colour=GREY, after=0)
    data = ["%s synthetic columns drawn from Panjabi's published means and spreads; every %s of every column "
            "scored, so most rates below are over tens of thousands of fields."
            % (count(N.columns), "field" if cns else "place")]
    if cns:
        deck.text(s, "question", x, 2.5, w, 1.15,
                  ["A view shows a few consecutive vertebral bodies, often with no C2 or sacrum to count from."], size=20)
        deck.text(s, "unknown", x, 3.75, w, 0.85, ["Seven measures per body. Which levels are in view?"], size=20,
                  bold=True, colour=NAVY)
        deck.text(s, "data", x, 4.8, w, 1.0, data, size=14, colour=GREY)
    else:
        deck.text(s, "question", x, 2.5, w, 1.15,
                  ["A view shows a few consecutive vertebral bodies, often with no C2 or sacrum to count from. "
                   "Which levels are they?"], size=20)
        deck.equation(s, "problem", x, 3.75, 20, max_w=w)
        deck.text(s, "unknown", x, 4.3, w, 0.95,
                  [[("Seven measures per body. Body ", {}), ("i", I), (" of the field is level ", {}),
                    ("k", I), (" + ", {}), ("i", I), (" \u2212 1, with C3 = 1 and L5 = %d, so the place " % N.levels, {}),
                    ("k", I), (" names every body in it.", {})]], size=16)
        deck.text(s, "data", x, 5.35, w, 1.0, data, size=14, colour=GREY)
    deck.footer(s)
    e1, e3 = N.exact(1), N.exact(3)
    if cns:
        deck.notes(s, (
            "Counting from C2 or the sacrum is how we name a level, and in the operating room the end we count from "
            "is often out of the picture. So the question is narrow: from the bodies we can see, and seven standard "
            "measures of each body, can we say which levels they are? Everything here is synthetic: %s columns drawn "
            "from Panjabi's published morphometry, so every answer is about the anatomy those tables describe, not "
            "yet about our patients." % count(N.columns)))
    else:
        deck.notes(s, (
            "The task: given the measurements of n consecutive vertebral bodies, seven per body, recover which levels "
            "they are. Because the bodies are consecutive, the whole labelling is one discrete latent variable, the "
            "field's place k: body i is level k + i - 1, with C3 = 1 and L5 = %d, so k takes %d minus n values. Data "
            "come from a generative model fitted to nothing: Panjabi's published per-level means and spreads, with a "
            "column-size factor, a within-column correlation and measurement noise. We score %s synthetic columns, "
            "every place of every column, and ask three questions: what is the best any function can do, can it say "
            "when it is right, and what can a single body say for certain. Headline numbers to come, on these "
            "synthetic columns: %s named exactly from one body, %s from three."
            % (N.levels, N.levels + 1, count(N.columns), pct(e1), pct(e3))))


def slide_bayes(deck: Deck, N: Numbers, cns: bool):
    if cns:
        return slide_more_bodies(deck, N)
    s = deck.slide("Bayes' rule over every place is the ceiling")
    x, w = 0.6, 5.75
    y = 1.5 + deck.equation(s, "bayes", x, 1.5, 24, max_w=w)[1]
    deck.text(s, "in words", x, y + 0.05, w, 0.42,
              [[("posterior", {"colour": BLUE, "bold": True}), ("  \u221d  ", {}), ("likelihood", {"bold": True}),
                ("  \u00d7  ", {}), ("prior", {"colour": GREY, "bold": True})]], size=18, after=0)
    y += 0.6
    y += deck.equation(s, "likelihood", x, y, 18, max_w=w)[1] + 0.08
    y += deck.equation(s, "covariance", x, y, 16, max_w=w)[1] + 0.08
    y += deck.equation(s, "size", x, y, 15, max_w=w)[1]
    deck.text(s, "how", x, y + 0.12, w, 6.8 - y - 0.12, [
        [("Every one of the %d\u00a0\u2212\u00a0" % (N.levels + 1), {}), ("n", I), (" places is scored under the prior ", {}), ("p", I),
         ("(", {}), ("k", I), (")\u00a0=\u00a01/(%d\u00a0\u2212\u00a0" % (N.levels + 1), {}), ("n", I), ("); nothing is fitted.", {})],
        [("\u03c3", I), (": the table's spreads at levels ", {}), ("k", I), (" to ", {}), ("k", I), (" + ", {}), ("n", I),
         (" \u2212 1, scaled by the column's size ", {}), ("s", I),
         ("; the covariance is inverted in closed form and ", {}), ("log s", I),
         (" is integrated by Gauss\u2013Hermite quadrature on %d nodes." % N.nodes, {})],
        "So it is the Bayes rule for this generator: no formula or network on the same measures names more "
        "fields correctly on average."], size=14, after=6)
    deck.figure(s, "exact_slide_posterior", 6.45, 1.5)
    deck.footer(s)
    deck.notes(s, (
        "Exact inference by enumeration. For a field of n bodies there are %d minus n places, each with prior 1 over "
        "%d minus n. Given the column's size s, each measure's n values are jointly Gaussian with mean s times the "
        "table means at levels k to k plus n minus 1, and covariance s squared rho squared sigma sigma-transpose plus "
        "a diagonal of s squared times 1 minus rho squared times sigma squared, plus the error variance, with rho %g "
        "and %g mm of error; the rank-one term means the likelihood needs no matrix inverse (Sherman-Morrison). "
        "Measures are independent given s; log s is Gaussian with SD %g and is integrated by %d-node Gauss-Hermite "
        "quadrature, whose error we measured. Under the uniform prior the posterior is the normalised likelihood, "
        "shown right for one column chosen by a stated rule: of %s, the column whose posterior at T12 sits nearest "
        "the median at one, three and nine bodies together. Because this is the generator's own posterior it is the "
        "Bayes-optimal classifier: a network can match it in expectation, never beat it, which turns 'can a network "
        "do this' into 'how far is this ceiling from one'."
        % (N.levels + 1, N.levels + 1, N.rho, N.noise_mm, N.tau, N.nodes, count(N.columns))))


def slide_more_bodies(deck: Deck, N: Numbers):
    """The CNS slide 2: the result alone, no rule."""
    s = deck.slide("More bodies in view, more certainty about the level")
    x, w = 0.6, 5.65
    deck.text(s, "result", x, 1.55, w, 1.0,
              ["Fields centred on T12 in synthetic columns, named exactly:"], size=20)
    for i, (n, words) in enumerate(((1, "from one body"), (3, "from three bodies"), (9, "from nine bodies"))):
        deck.stat(s, "T12 %d" % n, x, 2.55 + 1.0 * i, w, pct(N.by_centre(n)), words, big_size=36, small_size=16,
                  h=0.95)
    deck.text(s, "ceiling", x, 5.65, w, 1.1,
              ["No formula or network on these measures names more fields correctly on average."], size=16)
    deck.figure(s, "exact_slide_posterior", 6.45, 1.5)
    deck.footer(s)
    deck.notes(s, (
        "Before looking, every level is equally likely. Take the synthetic fields centred on T12. From one body the "
        "answer spreads over the levels around the thoracolumbar junction, T11 to L2 on average, and T12 is named "
        "exactly in %s of these synthetic fields; from three bodies, %s; from nine, %s. The figure shows one typical "
        "synthetic column. The rule is exact for the anatomy the tables describe, so no formula or network on these "
        "measures names more synthetic fields correctly on average."
        % (pct(N.by_centre(1)), pct(N.by_centre(3)), pct(N.by_centre(9)))))


def slide_ceiling(deck: Deck, N: Numbers, cns: bool):
    e3, sure3, sure9 = N.exact(3), N.sure(3), N.sure(9)
    lin, net = N.classifier(3), N.challenger("network")
    sure = XS.SURE[deck.v.name]
    s = deck.slide("Three bodies name %s; when sure, %s are right" % (pct(e3), pct(sure3["precision"], 1)))
    first = "of three-body fields named exactly (a linear classifier: %s; a fitted network: %s)" % (pct(lin, 1), pct(net, 1))
    second = ("of three-body names given at %s right (%s wrong of %s; named for %s of fields)"
              % (sure, count(sure3["errors"]), count(sure3["n_answered"]), pct(sure3["answered"])))
    if cns:
        low, w0, w1 = N.least_sure_precision()
        deck.figure(s, "exact_slide_ceiling", 0.6, 1.45)
        x, w = 8.4, 4.3
        deck.stat(s, "three bodies", x, 1.5, w, pct(e3), first, h=1.35)
        deck.stat(s, "precision", x, 2.9, w, pct(sure3["precision"], 1), second, h=1.35)
        deck.stat(s, "nine bodies", x, 4.3, w, pct(N.exact(9)), "named exactly from nine bodies", h=0.9)
        deck.text(s, "result", x, 5.3, w, 1.4,
                  ["Named only at 99%% confidence, it was right %s or more of the time at every number of bodies "
                   "in view, from %d to %d, on synthetic columns." % (pct(low, 1), w0, w1)], size=14)
        deck.footer(s)
        deck.notes(s, (
            "On synthetic columns drawn from the tables, with one body in view the rule names %s of fields exactly; "
            "with three, %s; with nine, %s. On the same synthetic measures a linear classifier reaches %s from three "
            "bodies and a fitted network %s. The rule also says how sure it is. On these synthetic columns, when it "
            "says 99 per cent it is right %s of the time, %s wrong of %s names, and it gives a name for %s of "
            "three-body fields; for the rest it says it is not sure."
            % (pct(N.exact(1)), pct(e3), pct(N.exact(9)), pct(lin, 1), pct(net, 1), pct(sure3["precision"], 1),
               count(sure3["errors"]), count(sure3["n_answered"]), pct(sure3["answered"]))))
        return
    deck.figure(s, "exact_slide_ceiling", 0.6, 1.45)
    deck.figure(s, "exact_slide_calibration", 6.25, 1.5)
    x, w = 10.05, 2.68
    deck.stat(s, "three bodies", x, 1.5, w, pct(e3), first, small_size=13, h=1.45)
    deck.stat(s, "precision", x, 3.0, w, pct(sure3["precision"], 1), second, small_size=13, h=1.5)
    deck.stat(s, "nine bodies", x, 4.55, w, pct(N.exact(9)), "named exactly from nine bodies", small_size=13, h=0.95)
    deck.text(s, "guarantee", 0.6, 5.72, 5.5, 1.05,
              [[("Answer only when ", {}), ("p\u0302", I),
                (" reaches 0.99, and the name is right at least 99% of the times it is given, in expectation, while "
                 "the generator is the truth.", {})]], size=14)
    _, h = deck.equation(s, "selective", 6.25, 5.68, 17, max_w=6.48)
    deck.equation(s, "estimate", 6.25, 5.68 + h + 0.08, 15, max_w=6.48)
    deck.footer(s)
    lo, hi = N.gap_range()
    top1, top3, top9 = N.top_bin(1), N.top_bin(3), N.top_bin(9)
    none = " and the top bins at one and three bodies counted none" if top1["errors"] == top3["errors"] == 0 else ""
    deck.notes(s, (
        "Left: exact-match accuracy of the Bayes rule against the number of bodies, %s at one, %s at three, %s at "
        "nine; the crosses are a linear classifier fitted on separate synthetic columns, %.1f to %.1f points below on "
        "the same first %s columns. On three bodies a fitted network reaches %s and the best fitted rule %s, against "
        "the rule's %s on the same fields. Dashed: the share answered under selective prediction at posterior 0.99, %s "
        "at three bodies and %s at nine. Right: calibration, error predicted against error counted, on the diagonal to "
        "about 1 in 1,000; at 1 in 10,000 the agreement rests on %d errors counted in %s nine-body fields with "
        "posterior above %g,%s. The guarantee is one line: among answers with max posterior "
        "at least 1 minus epsilon, accuracy equals the mean posterior, hence at least 1 minus epsilon, in expectation "
        "under the generator; it holds within any group chosen by the measurements, not within groups chosen by the "
        "truth. Observed: %s at three bodies (%s wrong of %s), %s at nine (%s wrong of %s)."
        % (pct(N.exact(1), 1), pct(e3, 1), pct(N.exact(9), 1), lo, hi, count(N.first_columns()), pct(net, 1),
           pct(N.challenger("qda_raw"), 1), pct(N.challenger("exact"), 1), pct(sure3["answered"], 1),
           pct(sure9["answered"], 1), top9["errors"], count(top9["n"]), top9["lo"], none,
           pct(sure3["precision"], 2), count(sure3["errors"]), count(sure3["n_answered"]), pct(sure9["precision"], 2),
           count(sure9["errors"]), count(sure9["n_answered"]))))


def slide_claims(deck: Deck, N: Numbers, cns: bool):
    s = deck.slide("One body tells its region and its side of C7/T1, rarely of T12/L1")
    deck.figure(s, "exact_slide_claims", 0.6, 1.45)
    ct, tl, tl9, reg, far = N.cut(1, "C7/T1"), N.cut(1, "T12/L1"), N.cut(9, "T12/L1"), N.regions(1), N.region_far(1)
    sure = XS.SURE[deck.v.name]
    overall = "T12/L1 from one body: %s, %s wrong of %s" % (pct(tl["precision"], 2), count(tl["errors"]),
                                                            count(tl["n_decided"]))
    x, w = 6.25, 6.48
    if cns:
        deck.text(s, "claim rule", x, 1.55, w, 1.4,
                  ["Claims made at 99%% confidence are right at least 99%% of the time over all fields (%s). The "
                   "rates below are for bodies within two levels of a border, chosen by their true level, where a "
                   "claim can fall short." % overall], size=15)
        y = 3.05
    else:
        cw, ch = deck.equation(s, "claim", x, 1.55, 18, max_w=w)
        mw, mh = deck.equation(s, "middle", x + cw + 0.45, 1.55 + (ch - 0.3) / 2, 16, max_w=w - cw - 0.45)
        deck.text(s, "claim rule", x, 1.55 + ch + 0.1, w, 1.25,
                  [[("\u2113", I), (" is the middle body. Claim ", {}), ("\u2113 \u2265 c", I),
                    (" when this reaches 0.99 and ", {}), ("\u2113 < c", I),
                    (" when it falls to 0.01: right at least 99%% of the times made, over all fields (%s). Below, "
                     "bodies within two levels of a border, a group chosen by the true level, where it need not hold."
                     % overall, {})]], size=14)
        y = 1.55 + ch + 1.4
    sw, gap, sh = 3.1, 0.28, 1.45
    deck.stat(s, "C7/T1", x, y, sw, pct(ct["decided_within_two"]),
              "of single bodies within two levels of C7/T1 say which side; right %s (%s wrong of %s)"
              % (pct(ct["precision_within_two"], 1), count(ct["errors_within_two"]), count(ct["n_decided_within_two"])),
              h=sh)
    deck.stat(s, "T12/L1", x + sw + gap, y, sw, pct(tl["decided_within_two"]),
              [("of single bodies within two levels of T12/L1 say which side; right ", {}),
               ("%s (%s wrong of %s), short of 99%%" % (pct(tl["precision_within_two"], 1), count(tl["errors_within_two"]),
                                                        count(tl["n_decided_within_two"])), {"colour": RED, "bold": True})],
              h=sh)
    deck.stat(s, "region", x, y + sh + 0.1, sw, pct(reg["decided"]),
              "of single bodies say their region; right %s (%s wrong of %s), only %s of them far from both junctions"
              % (pct(1 - reg["errors"] / reg["n_decided"], 1), count(reg["errors"]), count(reg["n_decided"]),
                 count(far["wrong"])), h=sh)
    deck.stat(s, "nine bodies", x + sw + gap, y + sh + 0.1, sw, pct(tl9["decided_within_two"]),
              "of fields with middle body T11 to L1 say which side of T12/L1 from nine bodies (%s wrong of %s)"
              % (count(tl9["errors_within_two"]), count(tl9["n_decided_within_two"])), h=sh)
    deck.footer(s)
    if cns:
        deck.notes(s, (
            "Short of a name, what can one body say for certain? On the synthetic columns, two things. Its region, "
            "cervical, thoracic or lumbar, for %s of single synthetic bodies, wrong %s times in %s. And which side of "
            "the cervicothoracic junction it is on: near C7/T1 a single synthetic body says so %s of the time, right "
            "%s. Over all synthetic fields, a claim made at 99%% confidence is right at least 99%% of the time; at "
            "T12/L1, %s. Near the thoracolumbar junction a single synthetic body commits only %s of the time and is "
            "right %s, too weak to rely on even in this synthetic setting. Nine bodies say which side for %s of "
            "synthetic fields whose middle body is T11 to L1."
            % (pct(reg["decided"]), count(reg["errors"]), count(reg["n_decided"]), pct(ct["decided_within_two"]),
               pct(ct["precision_within_two"], 1), pct(tl["precision"], 2), pct(tl["decided_within_two"]),
               pct(tl["precision_within_two"], 1), pct(tl9["decided_within_two"]))))
    else:
        deck.notes(s, (
            "Weaker statements inherit the guarantee. With l = k + floor((n - 1)/2) the middle body, P(l at or "
            "caudal to c) is a sum of the same posterior over k, so a claim made at 0.99 on either side is right at "
            "least 99 per cent of the time over all fields, in expectation: at T12/L1 from one body, %s, %s wrong of "
            "%s. The figure conditions on the truth, fields whose middle body lies within two levels of each border, "
            "where no guarantee applies, so these are measurements. One body: the side of C7/T1 for %s, precision %s; "
            "the side of T12/L1 for only %s, precision %s, below the nominal level because the group is chosen by the "
            "truth; the region for %s of all single bodies, %s wrong of %s, %s of them far from both junctions. Nine "
            "bodies settle T12/L1 for %s of fields whose middle body is T11 to L1, the levels a nine-body field can "
            "centre on near it."
            % (pct(tl["precision"], 2), count(tl["errors"]), count(tl["n_decided"]), pct(ct["decided_within_two"], 1),
               pct(ct["precision_within_two"], 2), pct(tl["decided_within_two"], 1), pct(tl["precision_within_two"], 2),
               pct(reg["decided"], 1), count(reg["errors"]), count(reg["n_decided"]), count(far["wrong"]),
               pct(tl9["decided_within_two"], 1))))


def slide_robust(deck: Deck, N: Numbers, cns: bool):
    s = deck.slide("The weak point is the table, not the mathematics")
    deck.figure(s, "exact_slide_robust", 0.6, 1.5)
    spec_n, spec_h = N.world("table_specimens", "nominal"), N.world("table_specimens", "honest")
    spec_p = N.world("table_specimens", "predictive")
    err3_n, err3_h = N.world("noise_mm=3", "nominal"), N.world("noise_mm=3", "honest")
    lumbar = N.called("honest", "lumbar")
    sure3 = N.sure(3)
    sure = XS.SURE[deck.v.name]
    x, w = 8.2, 4.53
    deck.stat(s, "twelve spines", x, 1.5, w, pct(spec_n["precision_099"], 1),
              "of three-body names at %s right when the rule's table is estimated from %d spines, as Panjabi's "
              "tables are" % (sure, N.table_n), colour=RED, h=1.3)
    corrected = ("with a corrected rule, its one setting chosen on a separate pilot" if cns else
                 "with spreads widened and %g mm of error (the %g mm chosen on a separate pilot)" % (N.honest_mm, N.honest_mm))
    deck.stat(s, "widened", x, 2.9, w, pct(spec_h["precision_099"], 1),
              "%s, but %s of fields named, not %s, and lumbar names right %s (%s wrong of %s)"
              % (corrected, pct(spec_h["answered_099"], 1), pct(spec_n["answered_099"], 1), pct(lumbar["precision"], 1),
                 count(lumbar["errors"]), count(lumbar["n_answered"])), colour=INK, h=1.6)
    if not cns:
        deck.equation(s, "honest", x, 4.65, 16, max_w=w)
    nxt = ("test the rule on real full-spine CT." if cns else
           "measure the table on full-spine CTs counted from C2, and carry a real series' prior over places.")
    deck.text(s, "next", 0.6, 5.75, 12.13, 0.9, [[("Next: ", {"bold": True}), (nxt, {})]], size=16)
    deck.footer(s)
    if cns:
        deck.notes(s, (
            "Where can this go wrong? Not in the arithmetic: it is exact for the anatomy it assumes. The weak point is "
            "the anatomy table itself. Panjabi measured %d spines; a rule built on a table that small, scored on "
            "synthetic spines drawn from a different table, is right %s of the time when it says 99, not 99. A "
            "corrected rule, its one setting chosen on a separate pilot, is right %s on those synthetic spines, but it "
            "then names only %s of synthetic fields, and its lumbar names are right %s. So the next step is to test "
            "the rule on real full-spine CT."
            % (N.table_n, pct(spec_n["precision_099"], 1), pct(spec_h["precision_099"], 1),
               pct(spec_h["answered_099"], 1), pct(lumbar["precision"], 1))))
    else:
        cols, worlds, per = N.robust_columns
        deck.notes(s, (
            "Robustness, three bodies. On the main scoring columns the printed-table rule is right %s at 0.99, naming "
            "%s of fields. The rows here are scored on separate synthetic columns, %d per error world and %d worlds of "
            "%d columns for a table world. When the rule's table is itself estimated from %d specimens, the setting "
            "Panjabi's tables come from, precision drops to %s; with 3 mm of measurement error, %s. Inflating each "
            "spread by the root of 1 plus 1 over %d and assuming %g mm of error, the %g mm chosen on a separate "
            "pilot, brings precision back to %s in the twelve-spine worlds, at the cost of coverage, %s answered "
            "instead of %s, and unevenly, lumbar names %s. With 3 mm of error on the printed table it reaches only %s, "
            "so it fails there. The predictive rule, with nothing tuned, reaches %s at %s named in the same "
            "twelve-spine worlds. The posterior is only as good as the table, so the open problems are a table "
            "measured on full-spine CT counted from C2 and a prior from a real series."
            % (pct(sure3["precision"], 2), pct(sure3["answered"], 1), cols, worlds, per, N.table_n,
               pct(spec_n["precision_099"], 1), pct(err3_n["precision_099"], 1), N.table_n, N.honest_mm, N.honest_mm,
               pct(spec_h["precision_099"], 1), pct(spec_h["answered_099"], 1), pct(spec_n["answered_099"], 1),
               pct(lumbar["precision"], 1), pct(err3_h["precision_099"], 1), pct(spec_p["precision_099"], 1),
               pct(spec_p["answered_099"], 1))))


SLIDES = (slide_problem, slide_bayes, slide_ceiling, slide_claims, slide_robust)


def build(venue: str, out: str, record: str = RECORD, figures: Optional[str] = None, refresh: bool = False,
          eqdir: Optional[str] = None, public: bool = False) -> str:
    """One deck. `figures` is the NeurIPS set's folder, the CNS set being its cns/ folder (default:
    <record>/slide-figures). The equation pictures go beside the deck (`<outdir>/equations/`) unless
    `eqdir` says where. A saved deck that holds a refused word is deleted and the build stops."""
    figures = XS.venue_dir(figures or os.path.join(record, "slide-figures"), venue)
    eqdir = eqdir or os.path.join(os.path.dirname(os.path.abspath(out)), "equations")
    N = Numbers(record)
    deck = Deck(Venue(venue), figures, refresh, eqdir, equations(N), public)
    for make in SLIDES:
        make(deck, N, venue == "cns")
    deck.check()
    deck.properties()
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    deck.prs.save(out)
    rewrite_app_xml(out, deck.count)
    bad = refusals(out, venue == "cns")
    if bad:
        os.remove(out)
        raise SystemExit("refusing the %s deck: %s" % (venue, "; ".join("%r in %s" % (w, p) for p, w in bad[:12])))
    return out


def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--venue", choices=("neurips", "cns", "both"), default="both")
    ap.add_argument("--outdir", default=OUT)
    ap.add_argument("--record", default=RECORD)
    ap.add_argument("--figures", default=None,
                    help="the NeurIPS slide figures, the CNS set in its cns/ folder (default: <record>/slide-figures)")
    ap.add_argument("--refresh-equations", action="store_true", dest="refresh")
    ap.add_argument("--public", action="store_true",
                    help="leave out the footer's confidentiality marking; for use only after the provisional is filed "
                         "and the TTO has read the deck")
    a = ap.parse_args(argv)
    for v in (("neurips", "cns") if a.venue == "both" else (a.venue,)):
        print("wrote", build(v, os.path.join(a.outdir, "exact_slides_%s.pptx" % v), a.record, a.figures, a.refresh,
                             public=a.public))
    return 0


if __name__ == "__main__":
    sys.exit(main())
