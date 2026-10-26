"""Six slides on the CNS 2026 speaker template: the junction named from Panjabi's tables.

    python slides/make_junction_slides.py --template <the deck>.pptx --out out/junction_slides.pptx

The template is the owner's own deck, which this repository does not hold; every slide is built
on its layout '1_Title and Content Vertical' with the deck's own title box. Every number is read
from the two recorded runs' JSON (records/panjabi-2026-10-26/, three measures and four), never
typed; the figures are junction_figures.py's, built from the three-measure record into
<out dir>/figures unless --figures names a folder already built. No word of any withheld method
appears on a slide, in a table or in a note: the build asserts it and refuses the deck otherwise.
"""
from __future__ import annotations

import argparse
import json
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, "bordercall"))
RECORDS = os.path.join(ROOT, "records", "panjabi-2026-10-26")
NAVY, GREY = RGBColor(0x19, 0x1D, 0x63), RGBColor(0x59, 0x59, 0x59)
FORBIDDEN = ("schwingnet", "c-arm", "sweep", "proprietary", "exposure", "bordercall", "fluoro")
X0, W = Emu(1750000), Emu(10150000)
MEASURE = {"endplate": "end-plate width", "h_post": "posterior body height", "canal_w": "canal width", "canal_ap": "canal depth"}
PRETTY = {"typical": "a typical column", "four_lumbar": "four lumbar bodies", "six_lumbar": "six lumbar bodies",
          "t13": "a thirteenth rib-bearing body", "lumbar_rib": "a rib on L1", "no_rib_t12": "no rib on T12"}
CITATION = ("Synthetic columns drawn from Panjabi's published means and spreads (Spine 1991;16:861, 888; "
            "1992;17:299); no patient data. Retrospective modelling, 2026.")


def pct(k, n):
    return 100.0 * k / n


def rate(d, key, j, cls):
    r = d["by_phenotype"]["%s/profile" % j][cls]
    return pct(r[key], r["called"]), r[key], r["called"]


def width_rate(d, j, w, key):
    r = d["by_width"]["%s/profile" % j][str(w)]
    return pct(r[key], r["called"]), r[key], r["called"]


def twin(d, kind, j):
    t = d["twins"]["profile"][kind][j]
    return pct(t["same_call"], t["n"]), t["same_call"], t["n"]


def blend_rates(d, key):
    """Every blend row's rate under one key, as (phenotype, blend, per cent)."""
    return [(r["phenotype"], r["blend"], pct(r[key], r["called"])) for r in d["blend_sweep"]["profile"]]


class Deck:
    def __init__(self, template: str, figures: str):
        self.prs = Presentation(template)
        self.figures = figures
        # drop every slide of the deck, keep the template's layouts
        lst = self.prs.slides._sldIdLst
        for sid in list(lst):
            self.prs.part.drop_rel(sid.rId)
            lst.remove(sid)
        self.layout = self.prs.slide_layouts[2]

    def add_text(self, slide, x, y, w, h, lines, size=13, color=NAVY, align=PP_ALIGN.LEFT, space_after=4):
        tb = slide.shapes.add_textbox(x, y, w, max(h, Emu(200000)))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        for i, line in enumerate(lines):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            para.space_after = Pt(space_after)
            run = para.add_run()
            run.text = line
            run.font.size = Pt(size)
            run.font.color.rgb = color
        return tb

    def blank(self, title: str):
        """A slide on the content layout with the deck's own title box and no other placeholder:
        the wide box at the top, anchored at its foot so a second line grows upward and never
        into the figure, 30 pt bold navy, as the template's own slides set it."""
        s = self.prs.slides.add_slide(self.layout)
        t = s.shapes.title
        t.left, t.top, t.width, t.height = Emu(1750000), Emu(300000), Emu(10150000), Emu(1000000)
        t.text = title
        t.text_frame.word_wrap = True
        t.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
        for para in t.text_frame.paragraphs:
            for run in para.runs:
                run.font.size = Pt(30)
                run.font.bold = True
                run.font.color.rgb = NAVY
        for ph in list(s.placeholders):
            if ph.placeholder_format.idx != 0:
                ph._element.getparent().remove(ph._element)
        return s

    def with_figure(self, title, figure, lines, caption, notes):
        s = self.blank(title)
        pic = s.shapes.add_picture(os.path.join(self.figures, figure), X0, Emu(1300000), width=W)
        if pic.height > Emu(3650000):
            scale = 3650000 / pic.height
            pic.height = Emu(3650000)
            pic.width = Emu(int(pic.width * scale))
            pic.left = Emu(int(X0 + (W - pic.width) / 2))
        y = Emu(pic.top + pic.height + 60000)
        self.add_text(s, X0, y, W, Emu(1000000), lines, size=13)
        self.add_text(s, X0, Emu(6250000), W, Emu(300000), [caption], size=9, color=GREY, space_after=0)
        s.notes_slide.notes_text_frame.text = notes
        return s


def build(template: str, out: str, json3: str, json4: str, figures: str) -> str:
    with open(json3, encoding="utf-8") as fh:
        J3 = json.load(fh)
    with open(json4, encoding="utf-8") as fh:
        J4 = json.load(fh)
    if figures is None:
        import junction_figures
        figures = os.path.join(os.path.dirname(os.path.abspath(out)) or ".", "figures")
        junction_figures.run(json3, figures)
    deck = Deck(template, figures)
    m3 = ", ".join(MEASURE[m] for m in J3["measures"])
    m4 = ", ".join(MEASURE[m] for m in J4["measures"] if m not in J3["measures"])

    # 1 the profile
    deck.with_figure(
        "The junction, named from the shape profile alone",
        "junction_profile.png",
        ["Panjabi's tables give every level a mean and a spread for the end-plate, the body and the canal; drawn to their millimetres, the column is the shape a reader of CT sees",
         "The cervicothoracic step is large in every measure; the thoracolumbar step is small beside the spread of one level",
         "A window of bodies is named by scoring every naming the column allows against this profile, with no count from either end"],
        "Panjabi MM, et al. Spine 1991;16:861 and 888; Spine 1992;17:299. Published morphometry drawn to scale; the bars are one population SD. No patient data.",
        "The three published tables, cervical, thoracic and lumbar, as one column drawn to the means (A: each body a box of its end-plate width and posterior height, the canal width hatched over it, the disc spaces at one height since the tables give none) and as one profile per measure (B to D). The dispersion printed in the papers is the standard error; the bar drawn is the population SD, SEM times the square root of twelve. C2 has no upper end-plate row, so the profile starts at C3.")

    # 2 how it decides
    deck.with_figure(
        "How it decides: every naming scored, the gap explained",
        "junction_posterior.png",
        ["Every naming a column with 11 to 13 thoracic and 4 to 6 lumbar bodies allows is scored per measure; a free offset per measure cancels the patient's size and the measurement convention",
         "The posterior over the cut, averaged over the namings that share it under a flat prior, is the answer; the per-measure gap to the runner-up says why: at C7/T1 the canal width decides, at T12/L1 no measure decides strongly",
         "A margin below a chosen floor is an abstention, not a guess"],
        "One synthetic column from the tables, eight bodies around each junction; " + CITATION,
        "The template holds the profile for every structure a column can have: eleven, twelve or thirteen rib-bearing bodies above four, five or six lumbar ones. For a window of n bodies every naming such a column allows is scored: per measure, the log-likelihood of the window's values against the profile's means and spreads, with one free offset per measure, so a large patient and a small one, or two measurement conventions, score alike. The posterior over the cut, the number of bodies above the junction, is the average over the namings that put the cut in one place, under a flat prior over the cut, so a cut is not favoured for having more ways to be named. The explanation a reader gets is the gap, per measure, between the best naming and the runner-up. On this column the canal width carries the cervicothoracic call by itself; at the thoracolumbar junction the gaps are within one log-likelihood unit, which is why that junction is the hard one.")

    # 3 the width result
    w8c, w8ck, w8cn = width_rate(J3, "ct", 8, "exact")
    w8c1 = width_rate(J3, "ct", 8, "within_one")[0]
    w8t = width_rate(J3, "tl", 8, "exact")[0]
    w8t1 = width_rate(J3, "tl", 8, "within_one")[0]
    w16c, w16t = width_rate(J3, "ct", 16, "exact")[0], width_rate(J3, "tl", 16, "exact")[0]
    w4c, w4t = width_rate(J3, "ct", 4, "exact")[0], width_rate(J3, "tl", 4, "exact")[0]
    local = {j: pct(sum(v["exact"] for v in J3["by_phenotype"]["%s/local" % j].values()),
                    sum(v["called"] for v in J3["by_phenotype"]["%s/local" % j].values())) for j in ("ct", "tl")}
    deck.with_figure(
        "Result: C7/T1 named robustly, T12/L1 the hard junction",
        "junction_width.png",
        ["Eight bodies in view: C7/T1 named exactly in %.0f of 100 windows and within one level in %.0f; T12/L1 exactly in %.0f, within one in %.0f" % (w8c, w8c1, w8t, w8t1),
         "Sixteen bodies: %.0f and %.0f exactly. Four bodies: %.0f and %.0f" % (w16c, w16t, w4c, w4t),
         "Fitted on no CT: the model is read off the published tables, so every number is about the tables and the generator, not yet about a patient"],
        "Profile template, %d synthetic typical windows per width, measures %s; 95%% Wilson intervals. %s" % (w8cn, m3, CITATION),
        "Exact and within-one naming against the width of the window, for the profile template fitted on nothing but the tables. " +
        "; ".join("%s: %d bodies %d/%d exact" % (j.upper(), w, J3["by_width"]["%s/profile" % j][str(w)]["exact"], J3["by_width"]["%s/profile" % j][str(w)]["called"])
                  for j in ("ct", "tl") for w in (4, 5, 6, 8, 10, 12, 16)) +
        ". The ablation, a local rule on the ratios of each body to its neighbours, names C7/T1 in %.0f%% and T12/L1 in %.0f%% of windows of five to twelve bodies over every phenotype: the thoracolumbar step is too small for a local rule, which is why the whole profile is read." % (local["ct"], local["tl"]))

    # 4 twins and phenotypes
    t4c, t6c, t4t, t6t = twin(J3, "four_lumbar", "ct"), twin(J3, "six_lumbar", "ct"), twin(J3, "four_lumbar", "tl"), twin(J3, "six_lumbar", "tl")
    classes = ["typical", "four_lumbar", "six_lumbar", "t13", "lumbar_rib", "no_rib_t12"]
    ct_by = {c: rate(J3, "exact", "ct", c) for c in classes}
    tl_by = {c: rate(J3, "exact", "tl", c) for c in classes}
    ct_lo, ct_hi = min(v[0] for v in ct_by.values()), max(v[0] for v in ct_by.values())
    tl_worst = min(classes, key=lambda c: tl_by[c][0])
    tl_best = max(classes, key=lambda c: tl_by[c][0])
    deck.with_figure(
        "A lumbosacral anomaly does not move the junction call",
        "junction_twins.png",
        ["One column drawn once and named two ways, four or six lumbar bodies: the C7/T1 call is unchanged in %.1f and %.1f of 100 twins; T12/L1 in %.1f and %.1f" % (t4c[0], t6c[0], t4t[0], t6t[0]),
         "By phenotype, C7/T1 holds at %.0f to %.0f%% whatever the column; T12/L1 runs from %.0f%% with %s to %.0f%% with %s, %.0f%% on %s" % (ct_lo, ct_hi, tl_by[tl_worst][0], PRETTY[tl_worst], tl_by[tl_best][0], PRETTY[tl_best], tl_by["typical"][0], PRETTY["typical"]),
         "The lumbosacral count lies outside the junction's window; the transitional body lies inside it (next slide)"],
        "Profile template, %d synthetic windows per class, measures %s; 95%% Wilson intervals. %s" % (ct_by["typical"][2], m3, CITATION),
        "Twins: " + "; ".join("%s, %s: same call %d of %d, both right %d" % (k, j, J3["twins"]["profile"][k][j]["same_call"], J3["twins"]["profile"][k][j]["n"], J3["twins"]["profile"][k][j]["both_right"])
                              for k in ("four_lumbar", "six_lumbar") for j in ("ct", "tl")) +
        ". By phenotype, exact of %d: " % ct_by["typical"][2] + "; ".join("%s %s %d" % (j, c, J3["by_phenotype"]["%s/profile" % j][c]["exact"]) for j in ("ct", "tl") for c in classes) +
        ". A twin is the same column with the lumbosacral levels added or taken away; the profile template reads the window's position from C3, so the levels above the anomaly are identical in both twins and the call can move only where the window reaches the anomaly, the far end of a thoracolumbar window. The cervicothoracic window never holds the lumbosacral bodies.")

    # 5 the transitional body, two truths
    count3 = blend_rates(J3, "exact")
    within3 = blend_rates(J3, "within_one")
    ribs3 = blend_rates(J3, "exact_ribs")
    count4 = blend_rates(J4, "exact")
    lift = [c4[2] - c3[2] for c3, c4 in zip(count3, count4)]
    deck.with_figure(
        "The transitional body: two truths one level apart",
        "junction_blend.png",
        ["Blending the transitional body toward its neighbour's shape leaves T12/L1 naming by the count from C2 at %.0f to %.0f%% across the three phenotypes, and within one level at %.0f to %.0f%%" % (min(r[2] for r in count3), max(r[2] for r in count3), min(r[2] for r in within3), max(r[2] for r in within3)),
         "Scored against the ribs instead, the same calls are exact in only %.0f to %.0f%%: a rib-bearing L1 and a thirteenth thoracic body are one shape, and only the count from C2 tells them apart" % (min(r[2] for r in ribs3), max(r[2] for r in ribs3)),
         "That is the case a full-spine corpus counted from C2 meets, with the ribs read independently"],
        "Profile template, %d synthetic windows per point; solid, scored against the count; dashed, against the ribs. %s" % (J3["blend_sweep"]["profile"][0]["n"], CITATION),
        "Blend series, exact of %d against the count, then within one level, then against the ribs: " % J3["blend_sweep"]["profile"][0]["n"] +
        "; ".join("%s %.2f: %d, %d, %d" % (r["phenotype"], r["blend"], r["exact"], r["within_one"], r["exact_ribs"]) for r in J3["blend_sweep"]["profile"]) +
        ". The nineteenth body is drawn from its own region's shape (blend 0) toward its neighbour's (blend 1); the count names it L1 whatever its rib, the rib reading names the same column one level apart, and the two thoracolumbar phenotypes with a thirteenth rib are one anatomy drawn from the two ends of the blend. With the %s as a fourth measure the count-exact rate moves by %+.0f to %+.0f points." % (m4, min(lift), max(lift)))

    # 6 what is validated, what is not
    s = deck.blank("What is validated, what is not, and what comes next")
    names = [("typical", "Typical"), ("four_lumbar", "Four lumbar"), ("six_lumbar", "Six lumbar"),
             ("t13", "Thirteenth rib-bearing body"), ("lumbar_rib", "Rib on L1"), ("no_rib_t12", "No rib on T12")]
    rows = len(names) + 1
    tbl = s.shapes.add_table(rows, 5, X0, Emu(1350000), Emu(6300000), Emu(2900000)).table
    for c, h in enumerate(["Column", "C7/T1 exact", "C7/T1 within one", "T12/L1 exact", "T12/L1 within one"]):
        tbl.cell(0, c).text = h
    for r, (key, name) in enumerate(names, start=1):
        tbl.cell(r, 0).text = name
        for c, (j, k) in enumerate((("ct", "exact"), ("ct", "within_one"), ("tl", "exact"), ("tl", "within_one")), start=1):
            p_, kk, nn = rate(J3, k, j, key)
            tbl.cell(r, c).text = "%.0f%% (%d/%d)" % (p_, kk, nn)
    for r in range(rows):
        for c in range(5):
            cell = tbl.cell(r, c)
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(11)
                    run.font.bold = (r == 0)
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if r == 0 else NAVY
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (RGBColor(0xF2, 0xF4, 0xFA) if r % 2 else RGBColor(0xFF, 0xFF, 0xFF))
    deck.add_text(s, Emu(8250000), Emu(1350000), Emu(3650000), Emu(4500000), [
        "Validated: the lumbar measurements against published norms. Not validated: the same measures at cervical and thoracic levels, until each is checked against a published series there",
        "Every number on these slides is from synthetic columns drawn from Panjabi's tables; none is from a patient",
        "Next: full-spine CT counted from C2, every column's class with its denominator; where the shape and the count disagree, the ribs decide, never the model's own call, and the circular arm is shown beside it for scale",
        "The thoracolumbar junction is the one still owed"], size=12)
    deck.add_text(s, X0, Emu(4400000), Emu(6300000), Emu(600000), [
        "Named exactly and within one level, by phenotype, %d synthetic windows per class; the denominators are the windows with the junction in view. A thirteenth rib-bearing body and a rib on L1 are one anatomy drawn from the two ends of the blend." % rate(J3, "exact", "ct", "typical")[2]], size=9, color=GREY)
    deck.add_text(s, X0, Emu(6250000), W, Emu(300000), [CITATION + " Evidence status: modelling on published tables; no clinical use."], size=9, color=GREY, space_after=0)
    s.notes_slide.notes_text_frame.text = ("Only the lumbar measurements of the measurement code are validated against published norms (end-plate width, canal width, transverse-process span); the thoracic and cervical figures are a pipeline test until each measure is checked against a published series at those levels, which the Panjabi tables are there for. The real test is the corpus counted from C2, by phenotype with denominators, with the rib arm as the independent relabelling and the circular arm beside it for scale, never alone.")

    deck.prs.save(out)
    bad = withheld_words(out)
    if bad:
        os.remove(out)
        raise SystemExit("a withheld word is on the slides: %r" % bad)
    return out


def withheld_words(path):
    """The public cut's rule: no withheld word anywhere, on a slide, in a table or in a note."""
    bad = []
    for i, sl in enumerate(Presentation(path).slides, start=1):
        texts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
        texts += [c.text for sh in sl.shapes if sh.has_table for row in sh.table.rows for c in row.cells]
        texts.append(sl.notes_slide.notes_text_frame.text)
        for t in texts:
            for w in FORBIDDEN:
                if w in t.lower():
                    bad.append((i, w))
    return bad


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--template", required=True, help="the deck whose layouts and title box the slides take")
    ap.add_argument("--out", required=True, help="the deck to write")
    ap.add_argument("--json", default=os.path.join(RECORDS, "three-measures", "panjabi_junction.json"), help="the recorded run the slides read")
    ap.add_argument("--json-more", default=os.path.join(RECORDS, "four-measures", "panjabi_junction.json"), dest="json_more",
                    help="the run with one more measure, for the notes' comparison")
    ap.add_argument("--figures", default=None, help="a folder junction_figures already built from --json; built beside --out otherwise")
    a = ap.parse_args(argv)
    out = build(a.template, a.out, a.json, a.json_more, a.figures)
    print("wrote", out, "slides", len(Presentation(out).slides))
    return 0


if __name__ == "__main__":
    sys.exit(main())
